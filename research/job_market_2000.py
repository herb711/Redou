#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a reproducible CS-employment company sample: >=1000 Chongqing + >=1000 outside.

All rows retain their source, date and verification level.  The script prefers direct
job records, then official software/technology enterprise lists.  It never invents a
company name.
"""
from __future__ import annotations

import csv
import hashlib
import hmac
import json
import os
import re
import shutil
import subprocess
import sys
import time
import uuid
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Tuple
from urllib.parse import quote, urlencode, urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / "research_work"
OUT = ROOT / "research_output"
WORK.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/149 Safari/537.36"
SESSION = requests.Session()
SESSION.headers.update({"User-Agent": UA, "Accept-Language": "zh-CN,zh;q=0.9"})

TECH_TITLE_KW = [
    "软件", "开发", "程序", "算法", "数据", "测试", "运维", "网络", "系统", "信息化", "信息技术",
    "前端", "后端", "全栈", "java", "python", "c++", "c#", "php", "golang", "go开发", "android", "ios",
    "嵌入式", "单片机", "驱动", "硬件", "芯片", "集成电路", "通信", "安全", "数据库", "云计算", "大数据",
    "人工智能", "机器学习", "深度学习", "图像", "视觉", "nlp", "大模型", "web", "ui", "ue", "游戏",
    "自动化", "电子", "物联网", "plc", "mes", "gis", "仿真", "fpga", "arm", "linux", "鸿蒙", "harmony",
    "车载", "智能座舱", "机器人", "产品经理", "实施工程师", "技术支持", "erp", "sap", "bi工程师",
    "信息安全", "网络安全", "数字化", "计算机", "研发工程师", "研发", "应用工程师", "架构师"
]
CORE_SECTOR_KW = [
    "计算机软件", "计算机硬件", "it服务", "互联网-电子商务", "互联网/电子商务", "通信-电信",
    "电子技术-半导体-集成电路", "仪器仪表及工业自动化", "网络游戏", "软件/互联网开发",
    "信息传输", "软件和信息技术", "人工智能", "大数据", "物联网", "工业互联网"
]
LEGAL_SUFFIX_RE = re.compile(
    r"(?:有限公司|有限责任公司|股份有限公司|集团有限公司|集团股份有限公司|研究院有限公司|科技中心|研究院|公司)$"
)

# Verified official 100 software and information-service enterprises, Chongqing HRSS, 2023-06-19.
OFFICIAL_100_TEXT = r"""
中移物联网有限公司|南岸区（含重庆经开区）
中冶赛迪信息技术（重庆）有限公司|西部科学城重庆高新区
重庆西南集成电路设计有限责任公司|南岸区（含重庆经开区）
重庆宅喵科技有限公司|两江新区
重庆首讯科技股份有限公司|渝北区
重庆中联信息产业有限责任公司|两江新区
北斗星通智联科技有限责任公司|渝北区
中科创达（重庆）汽车科技有限公司|渝北区
途作林杰科技有限公司|渝中区
重庆长安汽车软件科技有限公司|渝北区
中煤科工重庆设计研究院（集团）有限公司|渝中区
重庆市通信建设有限公司|九龙坡区
重庆啄木鸟网络科技有限公司|两江新区
博瑞得科技有限公司|两江新区
航天新通科技有限公司|西部科学城重庆高新区
重庆航天信息有限公司|九龙坡区
重庆梅安森科技股份有限公司|九龙坡区
重庆传音通讯技术有限公司|渝北区
重庆桴之科科技发展有限公司|巴南区
重庆金山科技（集团）有限公司|渝北区
重庆维普资讯有限公司|两江新区
重庆南华中天信息技术有限公司|九龙坡区
马上消费金融股份有限公司|两江新区
中冶赛迪工程技术股份有限公司|渝中区
重庆锐云科技有限公司|两江新区
重庆市信息通信咨询设计院有限公司|九龙坡区
重庆四通都成科技发展有限公司|九龙坡区
重庆中科云从科技有限公司|两江新区
重庆市城投金卡信息产业（集团）股份有限公司|南岸区（含重庆经开区）
博拉网络股份有限公司|两江新区
东方中讯数字证书认证有限公司|南岸区（含重庆经开区）
东联信息技术有限公司|两江新区
光辉城市（重庆）科技有限公司|两江新区
广域铭岛数字科技有限公司|两江新区
空间视创（重庆）科技股份有限公司|两江新区
数字重庆信息技术服务有限公司|两江新区
同炎数智科技（重庆）有限公司|重庆高新区
新驱动重庆智能汽车有限公司|两江新区
重庆贝特计算机系统工程有限公司|渝中区
重庆比特数图科技有限公司|南岸区（含重庆经开区）
重庆不贰科技（集团）有限公司|渝中区
重庆赋比兴科技有限公司|两江新区
重庆格工自动化控制设备有限公司|渝北区
重庆工业大数据创新中心有限公司|北碚区
重庆观度科技股份有限公司|大渡口区
重庆灏瀚网络科技有限公司|南岸区（含重庆经开区）
重庆忽米网络科技有限公司|九龙坡区
重庆华源智禾科技有限公司|渝中区
重庆汇集源科技有限公司|渝中区
重庆慧都科技有限公司|九龙坡区
重庆机电智能制造有限公司|南岸区（含重庆经开区）
重庆精耕企业管理咨询有限公司|大渡口区
重庆可兰达科技有限公司|两江新区
重庆朗维机电技术有限公司|两江新区
重庆旅游云信息科技有限公司|两江新区
重庆梦之想科技有限责任公司|两江新区
重庆耐德自动化技术有限公司|两江新区
重庆鹏康大数据有限公司|两江新区
重庆七彩虹数码科技有限公司|九龙坡区
重庆赛迪奇智人工智能科技有限公司|渝中区
重庆尚优科技有限公司|两江新区
重庆市地矿测绘院有限公司|渝中区
重庆市勘察规划设计有限公司|两江新区
重庆天智慧启科技有限公司|两江新区
重庆物奇科技有限公司|渝北区
重庆新世杰电气股份有限公司|渝北区
重庆新思维信息技术有限公司|渝中区
重庆新致金服信息技术有限公司|渝北区
重庆巽诺科技有限公司|大渡口区
重庆扬升科技集团有限公司|渝北区
重庆云辑数字科技有限责任公司|南岸区（含重庆经开区）
重庆云链数智软件科技有限公司|重庆高新区
重庆云网科技股份有限公司|南岸区（含重庆经开区）
重庆正大华日软件有限公司|两江新区
重庆知行数联智能科技有限责任公司|两江新区
重庆中科摇橹船信息科技有限公司|两江新区
重庆众鸿科技有限公司|南岸区（含重庆经开区）
重庆众康云科技有限责任公司|两江新区
重庆卓智软件开发有限公司|渝北区
紫光南方云技术有限公司|两江新区
重庆都会信息科技有限公司|江北区
爱思科技（重庆）集团有限公司|南岸区（含重庆经开区）
奥莫软件有限公司|渝北区
东方微银科技股份有限公司|江北区
菲欧坦（重庆）数据科技有限公司|两江新区
固守远望（重庆）信息科技有限公司|南岸区（含重庆经开区）
海尔数字科技（重庆）有限公司|江北区
撼地数智（重庆）科技有限公司|两江新区
数字重庆大数据应用发展有限公司|两江新区
益模（重庆）智能制造研究院有限公司|沙坪坝区
谊风信息技术（重庆）有限公司|渝中区
中管坤达（重庆）工业互联网有限公司|江北区
中再云图技术有限公司|西部科学城重庆高新区
重庆爱车天下科技有限公司|渝中区
重庆傲雄在线信息技术有限公司|两江新区
重庆傲洋科技有限公司|渝中区
重庆巴陆科技有限公司|两江新区
重庆巴云科技有限公司|南岸区（含重庆经开区）
重庆柏锐特智能科技有限公司|江北区
重庆宝图科技发展有限公司|两江新区
""".strip()

SOURCE_LOG: List[Dict[str, str]] = []
JOBS: List[Dict[str, str]] = []
OFFICIAL_COMPANIES: List[Dict[str, str]] = []


def log_source(source: str, url: str, status: str, rows: int = 0, note: str = "") -> None:
    SOURCE_LOG.append({
        "source": source, "url": url, "status": status, "rows": str(rows), "note": note,
        "checked_at": datetime.now().isoformat(timespec="seconds")
    })


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u3000", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def norm_company(name: str) -> str:
    name = norm_text(name).replace("(", "（").replace(")", "）")
    name = name.replace("【", "").replace("】", "")
    name = re.sub(r"^[·•\-—\s]+|[·•\-—\s]+$", "", name)
    return name


def is_company_name(name: str) -> bool:
    n = norm_company(name)
    if len(n) < 4 or len(n) > 100:
        return False
    return bool(LEGAL_SUFFIX_RE.search(n)) and not any(x in n for x in ["岗位", "招聘", "职位", "名单", "附件"])


def contains_kw(text: str, keywords: List[str]) -> bool:
    t = norm_text(text).lower()
    return any(k.lower() in t for k in keywords)


def direct_tech_job(title: str, desc: str = "") -> bool:
    return contains_kw(title, TECH_TITLE_KW) or contains_kw((title + " " + desc)[:1600], TECH_TITLE_KW)


def city_from_location(loc: str) -> str:
    loc = norm_text(loc)
    if not loc:
        return ""
    for city in ["重庆", "北京", "上海", "深圳", "广州", "杭州", "成都", "西安", "武汉", "南京", "苏州", "合肥", "长沙", "天津", "郑州", "济南", "青岛", "厦门", "福州", "宁波", "无锡", "东莞", "佛山", "珠海", "大连", "沈阳", "长春", "哈尔滨", "昆明", "贵阳", "南宁", "南昌", "石家庄", "太原", "兰州", "乌鲁木齐", "海口"]:
        if city in loc:
            return city
    return re.split(r"[-·/ ]", loc)[0][:12]


def add_job(**kwargs: Any) -> None:
    row = {k: norm_text(v) for k, v in kwargs.items()}
    row["company_name"] = norm_company(row.get("company_name", ""))
    if not is_company_name(row["company_name"]):
        return
    row.setdefault("city", city_from_location(row.get("location", "")))
    row.setdefault("region_group", "重庆" if "重庆" in (row.get("location", "") + row.get("city", "")) else "外地")
    row.setdefault("direct_tech_job", "是" if direct_tech_job(row.get("job_title", ""), row.get("job_desc", "")) else "否")
    JOBS.append(row)


def clone_repo(url: str, dest: Path) -> bool:
    if dest.exists():
        shutil.rmtree(dest)
    try:
        subprocess.run(["git", "clone", "--depth", "1", url, str(dest)], check=True, timeout=600)
        return True
    except Exception as exc:
        log_source("GitHub repository", url, "failed", note=str(exc))
        return False


def field_value(block: str, label: str) -> str:
    m = re.search(rf"{re.escape(label)}[：:]\s*([^\n]+)", block)
    return norm_text(m.group(1)) if m else ""


def after_label(block: str, label: str) -> str:
    lines = [norm_text(x) for x in block.splitlines()]
    for i, line in enumerate(lines):
        if line.rstrip("：:") == label.rstrip("：:") and i + 1 < len(lines):
            return lines[i + 1]
        if line.startswith(label):
            return norm_text(line[len(label):].lstrip("：:"))
    return ""


def parse_andinsbing_repo(repo: Path) -> int:
    base = repo / "src" / "DataClean" / "bigdata"
    count = 0
    if not base.exists():
        return 0
    for path in base.glob("*.org"):
        stem = path.stem
        city_hint = stem.rsplit("_", 1)[-1] if "_" in stem else ""
        sector = stem.rsplit("_", 1)[0] if "_" in stem else stem
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            continue
        blocks = [x.strip() for x in text.split("\n**********\n") if x.strip()]
        for i in range(0, len(blocks) - 4, 5):
            head, _, meta, desc, company_info = blocks[i:i + 5]
            hlines = [norm_text(x) for x in head.splitlines() if norm_text(x)]
            if len(hlines) < 2:
                continue
            title, company = hlines[0], hlines[1]
            industry = after_label(company_info, "公司行业") or sector
            industry_related = contains_kw(industry + " " + sector, CORE_SECTOR_KW)
            tech = direct_tech_job(title, desc)
            if not (tech or industry_related):
                continue
            location = field_value(meta, "工作地点") or city_hint
            add_job(
                company_name=company, job_title=title, location=location, city=city_from_location(location) or city_hint,
                salary=field_value(meta, "职位月薪"), experience=field_value(meta, "工作经验"),
                education=field_value(meta, "最低学历"), job_category=field_value(meta, "职位类别"),
                company_size=after_label(company_info, "公司规模"), company_type=after_label(company_info, "公司性质"),
                industry=industry, job_desc=desc[:3000], source_platform="智联招聘历史公开数据集",
                source_date="约2018", source_url="https://github.com/andinsbing/Analysis-of-College-Graduates-Employment-Orientation",
                source_file=str(path.relative_to(repo)), verification_level="B-结构化招聘记录",
                relevance_basis="直接技术岗位" if tech else "企业行业相关"
            )
            count += 1
    log_source("智联招聘历史公开数据集（行业×城市）", "https://github.com/andinsbing/Analysis-of-College-Graduates-Employment-Orientation", "ok", count)
    return count


def parse_liuaichao_repo(repo: Path) -> int:
    data_dir = repo / "机器学习" / "智联招聘预测" / "data"
    count = 0
    if not data_dir.exists():
        return 0
    for path in data_dir.glob("*.csv"):
        try:
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
                reader = csv.DictReader(f)
                for r in reader:
                    company = r.get("公司名称", "")
                    title = r.get("职位名称", "")
                    job_type = r.get("工作类型", "")
                    if not direct_tech_job(title, job_type) and not contains_kw(path.stem, TECH_TITLE_KW):
                        continue
                    location = r.get("工作地点", "")
                    add_job(
                        company_name=company, job_title=title, location=location, city=city_from_location(location),
                        salary=r.get("职位月薪", ""), experience=r.get("工作经验", ""), education=r.get("工作要求", ""),
                        job_category=job_type, company_size=r.get("公司规模", ""), company_type=r.get("公司类型", ""),
                        industry=job_type, job_desc="", source_platform="智联招聘历史公开数据集",
                        source_date=r.get("发布时间", "2019-09"),
                        source_url="https://github.com/liuaichao/python-work",
                        source_file=str(path.relative_to(repo)), verification_level="B-结构化招聘记录",
                        relevance_basis="直接技术岗位"
                    )
                    count += 1
        except Exception as exc:
            log_source("智联CSV", str(path), "partial", note=str(exc))
    log_source("智联招聘历史公开CSV（岗位方向）", "https://github.com/liuaichao/python-work", "ok", count)
    return count


def walk_json(obj: Any) -> Iterator[Dict[str, Any]]:
    if isinstance(obj, dict):
        yield obj
        for value in obj.values():
            yield from walk_json(value)
    elif isinstance(obj, list):
        for value in obj:
            yield from walk_json(value)


def first_value(d: Dict[str, Any], keys: List[str]) -> str:
    for k in keys:
        v = d.get(k)
        if isinstance(v, (str, int, float)) and norm_text(v):
            return norm_text(v)
    return ""


def parse_boss_repo(repo: Path) -> int:
    count = 0
    for path in list(repo.rglob("*.json")) + list(repo.rglob("*.jsonl")):
        if path.stat().st_size > 100 * 1024 * 1024:
            continue
        try:
            raw = path.read_text(encoding="utf-8", errors="replace")
            data: Any
            if path.suffix == ".jsonl":
                data = [json.loads(x) for x in raw.splitlines() if x.strip().startswith("{")]
            else:
                data = json.loads(raw)
        except Exception:
            continue
        seen_local = set()
        for d in walk_json(data):
            company = first_value(d, ["companyName", "fullCompanyName", "brandName", "company", "company_name"])
            title = first_value(d, ["jobName", "positionName", "job_title", "title", "position"])
            if not company or not title or not direct_tech_job(title):
                continue
            key = (company, title, first_value(d, ["cityName", "jobAreaString", "city", "location", "area"]));
            if key in seen_local:
                continue
            seen_local.add(key)
            location = key[2]
            add_job(
                company_name=company, job_title=title, location=location, city=city_from_location(location),
                salary=first_value(d, ["salaryDesc", "salary", "provideSalaryString"]),
                experience=first_value(d, ["jobExperience", "experience", "workYearString"]),
                education=first_value(d, ["jobDegree", "degree", "degreeString"]),
                company_size=first_value(d, ["companyScale", "companySize", "companySizeString"]),
                company_type=first_value(d, ["companyType", "companyTypeString", "stageName"]),
                industry=first_value(d, ["industryName", "industry", "companyIndustry"]),
                source_platform="BOSS直聘公开数据集", source_date="数据集记录时间",
                source_url="https://github.com/RuShi-Aurora/BigData-BossZhiPin", source_file=str(path.relative_to(repo)),
                verification_level="B-结构化招聘记录", relevance_basis="直接技术岗位"
            )
            count += 1
    log_source("BOSS直聘公开招聘数据集", "https://github.com/RuShi-Aurora/BigData-BossZhiPin", "ok" if count else "no_rows", count)
    return count


def fetch_51job_current(max_pages: int = 6) -> int:
    """Best-effort current 51job no-auth search API. Failure is logged, never fabricated."""
    secret = b"abfc8f9dcf8c3f3d8aa294ac5f2cf2cc7767e5592590f39c3f503271dd68562b"
    keywords = ["Java", "软件开发", "嵌入式", "软件测试", "前端开发", "Python", "人工智能", "运维", "网络安全", "数据开发", "C++", "鸿蒙"]
    count = 0
    failures = 0
    for keyword in keywords:
        for page in range(1, max_pages + 1):
            timestamp = str(int(time.time()))
            params = [
                ("api_key", "51job"), ("timestamp", timestamp), ("keyword", keyword), ("searchType", "2"),
                ("function", ""), ("industry", ""), ("jobArea", "000000"), ("jobArea2", ""), ("landmark", ""),
                ("metro", ""), ("salary", ""), ("workYear", ""), ("degree", ""), ("companyType", ""),
                ("companySize", ""), ("jobType", ""), ("issueDate", ""), ("sortType", "0"),
                ("pageNum", str(page)), ("requestId", ""), ("pageSize", "50"), ("source", "1"),
                ("accountId", ""), ("pageCode", "sou|sou|soulb")
            ]
            query = urlencode(params, quote_via=quote, safe="|")
            path_query = "/open/noauth/search-pc?" + query
            sign = hmac.new(secret, path_query.encode(), hashlib.sha256).hexdigest()
            headers = {
                "Accept": "application/json, text/plain, */*", "From-Domain": "51job_web",
                "Origin": "https://we.51job.com", "Referer": "https://we.51job.com/", "sign": sign,
                "uuid": uuid.uuid4().hex, "User-Agent": UA
            }
            url = "https://cupid.51job.com" + path_query
            try:
                resp = SESSION.get(url, headers=headers, timeout=25)
                if resp.status_code != 200:
                    failures += 1
                    break
                payload = resp.json()
                items = (((payload.get("resultbody") or {}).get("job") or {}).get("items") or [])
                if not items:
                    break
                for item in items:
                    title = norm_text(item.get("jobName"))
                    if not direct_tech_job(title):
                        continue
                    location = norm_text(item.get("jobAreaString"))
                    add_job(
                        company_name=item.get("fullCompanyName") or item.get("companyName"), job_title=title,
                        location=location, city=city_from_location(location), salary=item.get("provideSalaryString"),
                        experience=item.get("workYearString"), education=item.get("degreeString"),
                        company_size=item.get("companySizeString"), company_type=item.get("companyTypeString"),
                        industry=item.get("industryType1Str") or item.get("industryType2Str"),
                        source_platform="51Job（前程无忧）", source_date=item.get("issueDateString") or "2026-07",
                        source_url=item.get("jobHref") or url, source_file="51job current API",
                        verification_level="A-当前招聘记录", relevance_basis="直接技术岗位"
                    )
                    count += 1
                time.sleep(0.35)
            except Exception:
                failures += 1
                break
    status = "ok" if count else "blocked_or_empty"
    log_source("51Job当前公开搜索接口", "https://we.51job.com/pc/search", status, count, f"failed keyword/page groups={failures}")
    return count


def official_100() -> int:
    url = "https://rlsbj.cq.gov.cn/zwxx_182/tzgg/202306/t20230619_12076811.html"
    for line in OFFICIAL_100_TEXT.splitlines():
        name, district = line.split("|", 1)
        OFFICIAL_COMPANIES.append({
            "company_name": norm_company(name), "region_group": "重庆", "city": "重庆", "district": district,
            "industry": "软件和信息服务", "source_platform": "重庆市人力社保局官方名单", "source_date": "2023-06-19",
            "source_url": url, "verification_level": "A-官方重点软件企业名单", "relevance_basis": "官方软件信息服务企业"
        })
    log_source("重庆市重点软件和信息服务企业名单（100家）", url, "ok", 100)
    return 100


def extract_company_rows_from_xlsx(path: Path, source_url: str) -> int:
    wb = load_workbook(path, read_only=True, data_only=True)
    count = 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        name_col = None
        district_col = None
        header_row = 0
        for ri, row in enumerate(rows[:15]):
            for ci, value in enumerate(row):
                text = norm_text(value)
                if "企业名称" in text or text == "单位名称":
                    name_col, header_row = ci, ri
                if any(k in text for k in ["所属区县", "区县", "所在区县"]):
                    district_col = ci
            if name_col is not None:
                break
        for row in rows[header_row + 1:]:
            candidates: List[Tuple[int, str]] = []
            if name_col is not None and name_col < len(row):
                candidates.append((name_col, norm_text(row[name_col])))
            else:
                candidates.extend((ci, norm_text(v)) for ci, v in enumerate(row))
            company = ""
            for _, value in candidates:
                if is_company_name(value):
                    company = value
                    break
            if not company:
                continue
            district = norm_text(row[district_col]) if district_col is not None and district_col < len(row) else ""
            OFFICIAL_COMPANIES.append({
                "company_name": norm_company(company), "region_group": "重庆", "city": "重庆", "district": district,
                "industry": "科技型企业", "source_platform": "重庆市科学技术局官方名单", "source_date": "2026-07-07",
                "source_url": source_url, "verification_level": "C-官方科技企业名单", "relevance_basis": "科技型企业补充样本"
            })
            count += 1
    return count


def fetch_official_1062() -> int:
    page_url = "https://kjj.cq.gov.cn/zwxx_176/tzgg/202607/t20260707_15805740.html"
    try:
        resp = SESSION.get(page_url, timeout=30)
        resp.raise_for_status()
        html = resp.text
        soup = BeautifulSoup(html, "html.parser")
        links: List[str] = []
        for a in soup.find_all("a", href=True):
            href = urljoin(page_url, a["href"])
            text = norm_text(a.get_text(" "))
            if any(ext in href.lower() for ext in [".xlsx", ".xls", ".docx", ".pdf"]) or "入库科技型企业名单" in text:
                links.append(href)
        for m in re.findall(r"[\"']([^\"']+\.(?:xlsx|xls|docx|pdf)(?:\?[^\"']*)?)[\"']", html, flags=re.I):
            links.append(urljoin(page_url, m))
        links = list(dict.fromkeys(links))
        for idx, link in enumerate(links):
            try:
                r = SESSION.get(link, timeout=60)
                r.raise_for_status()
                suffix = ".xlsx" if ("spreadsheet" in r.headers.get("content-type", "") or ".xlsx" in link.lower()) else Path(link.split("?", 1)[0]).suffix.lower()
                fpath = WORK / f"cq_official_1062_{idx}{suffix or '.bin'}"
                fpath.write_bytes(r.content)
                if suffix == ".xlsx":
                    n = extract_company_rows_from_xlsx(fpath, page_url)
                    if n >= 500:
                        log_source("重庆市2026年第六批科技型企业名单", page_url, "ok", n, f"attachment={link}")
                        return n
            except Exception:
                continue
        log_source("重庆市2026年第六批科技型企业名单", page_url, "attachment_not_parsed", 0, f"candidate_links={len(links)}")
    except Exception as exc:
        log_source("重庆市2026年第六批科技型企业名单", page_url, "failed", 0, str(exc))
    return 0


def quality_score(row: Dict[str, str]) -> int:
    level = row.get("verification_level", "")
    score = 0
    if level.startswith("A-当前"):
        score += 100
    elif level.startswith("A-官方重点"):
        score += 80
    elif level.startswith("B-"):
        score += 60
    elif level.startswith("C-"):
        score += 20
    if row.get("direct_tech_job") == "是" or row.get("relevance_basis") == "直接技术岗位":
        score += 30
    if row.get("salary"):
        score += 5
    if row.get("education"):
        score += 3
    return score


def merge_company_records() -> Tuple[List[Dict[str, str]], List[Dict[str, str]]]:
    grouped: Dict[str, List[Dict[str, str]]] = defaultdict(list)
    for job in JOBS:
        grouped[norm_company(job["company_name"])].append(job)
    for official in OFFICIAL_COMPANIES:
        grouped[norm_company(official["company_name"])].append(official)

    companies: List[Dict[str, str]] = []
    all_fields = [
        "company_name", "region_group", "city", "district", "company_type", "company_size", "industry",
        "job_title", "job_category", "salary", "education", "experience", "direct_tech_job", "relevance_basis",
        "source_platform", "source_date", "source_url", "source_file", "verification_level", "record_count"
    ]
    for company, rows in grouped.items():
        rows_sorted = sorted(rows, key=quality_score, reverse=True)
        best = rows_sorted[0]
        merged = {f: best.get(f, "") for f in all_fields}
        merged["company_name"] = company
        merged["record_count"] = str(len(rows))
        # If a job record has no district but an official record does, fill it.
        for f in ["district", "company_type", "company_size", "industry", "city"]:
            if not merged.get(f):
                merged[f] = next((r.get(f, "") for r in rows_sorted if r.get(f)), "")
        # Region is based on actual workplace where available, otherwise official location.
        is_cq = any(r.get("region_group") == "重庆" or "重庆" in (r.get("location", "") + r.get("city", "")) for r in rows)
        merged["region_group"] = "重庆" if is_cq else "外地"
        companies.append(merged)

    cq = sorted([x for x in companies if x["region_group"] == "重庆"], key=lambda x: (-quality_score(x), x["company_name"]))
    outside = sorted([x for x in companies if x["region_group"] == "外地"], key=lambda x: (-quality_score(x), x["company_name"]))
    return cq, outside


def write_csv(path: Path, rows: List[Dict[str, str]], fields: Optional[List[str]] = None) -> None:
    if fields is None:
        fields = sorted({k for r in rows for k in r.keys()}) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def main() -> None:
    print("[1/7] official lists")
    official_100()
    official_1062_count = fetch_official_1062()

    print("[2/7] clone structured recruitment repositories")
    repo_a = WORK / "andinsbing"
    repo_b = WORK / "liuaichao"
    repo_c = WORK / "bossdata"
    if clone_repo("https://github.com/andinsbing/Analysis-of-College-Graduates-Employment-Orientation.git", repo_a):
        parse_andinsbing_repo(repo_a)
    if clone_repo("https://github.com/liuaichao/python-work.git", repo_b):
        parse_liuaichao_repo(repo_b)
    if clone_repo("https://github.com/RuShi-Aurora/BigData-BossZhiPin.git", repo_c):
        parse_boss_repo(repo_c)

    print("[3/7] current 51job best-effort")
    current_51_count = fetch_51job_current()

    print("[4/7] deduplicate and rank")
    cq, outside = merge_company_records()
    print(f"unique Chongqing={len(cq)}, outside={len(outside)}, jobs={len(JOBS)}, official1062={official_1062_count}")
    if len(cq) < 1000 or len(outside) < 1000:
        raise RuntimeError(f"Target not met: Chongqing={len(cq)}, outside={len(outside)}")

    sample_cq = cq[:1000]
    sample_out = outside[:1000]
    sample = []
    for idx, row in enumerate(sample_cq + sample_out, 1):
        x = dict(row)
        x["sample_id"] = f"CSJOB-{idx:04d}"
        sample.append(x)

    print("[5/7] write data files")
    company_fields = [
        "sample_id", "company_name", "region_group", "city", "district", "company_type", "company_size", "industry",
        "job_title", "job_category", "salary", "education", "experience", "direct_tech_job", "relevance_basis",
        "source_platform", "source_date", "source_url", "source_file", "verification_level", "record_count"
    ]
    write_csv(OUT / "company_sample_2000.csv", sample, company_fields)
    write_csv(OUT / "chongqing_companies_1000.csv", [{**r, "sample_id": f"CQ-{i:04d}"} for i, r in enumerate(sample_cq, 1)], company_fields)
    write_csv(OUT / "outside_companies_1000.csv", [{**r, "sample_id": f"OUT-{i:04d}"} for i, r in enumerate(sample_out, 1)], company_fields)
    write_csv(OUT / "jobs_all.csv", JOBS)
    write_csv(OUT / "source_log.csv", SOURCE_LOG)

    platform_counts = Counter(r.get("source_platform", "") for r in sample)
    industry_counts = Counter(r.get("industry", "") for r in sample)
    city_counts = Counter(r.get("city", "") for r in sample_out)
    verification_counts = Counter(r.get("verification_level", "") for r in sample)
    direction_counts = Counter(r.get("job_category", "") or r.get("job_title", "") for r in sample if r.get("job_title"))
    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "target": {"chongqing": 1000, "outside": 1000, "total": 2000},
        "actual_unique_pool": {"chongqing": len(cq), "outside": len(outside)},
        "selected_sample": {"chongqing": len(sample_cq), "outside": len(sample_out), "total": len(sample)},
        "job_records_total": len(JOBS), "official_1062_parsed": official_1062_count,
        "current_51job_records": current_51_count,
        "sample_platform_counts": platform_counts.most_common(),
        "sample_verification_counts": verification_counts.most_common(),
        "outside_city_top20": city_counts.most_common(20),
        "industry_top30": industry_counts.most_common(30),
        "job_direction_top30": direction_counts.most_common(30),
        "limitations": [
            "Boss/51Job live pages may impose login, JavaScript or anti-bot restrictions; failed requests are logged rather than filled with invented data.",
            "Historical structured recruitment datasets are marked B and retain their original approximate dates.",
            "Official technology-enterprise rows without a directly matched job are marked C and rank below direct job records."
        ]
    }
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    readme = f"""# 计算机科学与技术专业就业企业2000家调研样本\n\n- 重庆样本：{len(sample_cq)} 家\n- 外地样本：{len(sample_out)} 家\n- 企业池：重庆 {len(cq)} 家、外地 {len(outside)} 家\n- 结构化岗位记录：{len(JOBS)} 条\n- 当前51Job成功记录：{current_51_count} 条\n- 重庆科技局第六批名单成功解析：{official_1062_count} 家\n\n## 核验等级\n- A：当前招聘记录或政府重点软件企业实名名单。\n- B：公开结构化招聘数据集中的真实岗位记录，保留历史日期。\n- C：政府科技型企业实名名单，尚未匹配到直接岗位，作为企业母库补充。\n\n所有公司名称均来自所列公开来源，未人工生成。\n"""
    (OUT / "README.md").write_text(readme, encoding="utf-8")
    print("[7/7] completed")


if __name__ == "__main__":
    main()
